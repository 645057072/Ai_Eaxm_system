# -*- coding: utf-8 -*-
"""题库文件解析：按文本块推断题型并生成草稿题目结构。"""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

# 题干最大字符数（与接口校验一致）
STEM_MAX_LEN = 2000
# 解析字段最大字符数（MySQL TEXT 约 64KB；UTF-8 下单字段不宜过长）
ANALYSIS_MAX_LEN = 16000


def _text_from_txt_csv(raw: bytes) -> str:
    return raw.decode("utf-8", errors="replace")


def _text_from_xlsx(raw: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.active
    lines: List[str] = []
    for row in ws.iter_rows(values_only=True):
        lines.append("\t".join("" if c is None else str(c).strip() for c in row))
    return "\n".join(lines)


def _text_from_xls(raw: bytes) -> str:
    """读取老版 Excel（.xls）。"""
    import xlrd

    book = xlrd.open_workbook(file_contents=raw)
    sh = book.sheet_by_index(0)
    lines: List[str] = []
    for r in range(sh.nrows):
        parts: List[str] = []
        for c in range(sh.ncols):
            v = sh.cell_value(r, c)
            parts.append("" if v is None else str(v).strip())
        lines.append("\t".join(parts))
    return "\n".join(lines)


def _text_from_docx(raw: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(raw))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _text_from_pdf(raw: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    parts: List[str] = []
    for page in reader.pages:
        t = page.extract_text()
        if t:
            parts.append(t)
    return "\n".join(parts)


def extract_plain_text(filename: str, raw: bytes) -> str:
    """从上传文件提取纯文本；图片等非文本返回空串，由上层生成占位题干。"""
    ext = Path(filename).suffix.lower()
    if ext in (".txt",):
        return _text_from_txt_csv(raw)
    if ext == ".csv":
        text = _text_from_txt_csv(raw)
        lines = []
        try:
            for row in csv.reader(io.StringIO(text)):
                if row:
                    lines.append(row[0].strip())
        except Exception:
            return text
        return "\n".join(lines)
    if ext == ".xlsx":
        return _text_from_xlsx(raw)
    if ext == ".xls":
        return _text_from_xls(raw)
    if ext == ".docx":
        return _text_from_docx(raw)
    if ext == ".pdf":
        return _text_from_pdf(raw)
    if ext in (".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"):
        return ""
    raise ValueError(f"不支持的扩展名: {ext}")


def _truncate_stem(s: str) -> str:
    s = s.strip()
    if len(s) <= STEM_MAX_LEN:
        return s
    return s[:STEM_MAX_LEN]


def _normalize_import_text(text: str) -> str:
    """导入前规范化换行，减轻 PDF 抽字粘连。"""
    t = text.replace("\r\n", "\n").replace("\r", "\n")
    # 「理论部分」与页码粘连：理论部分2 -> 理论部分\n2
    t = re.sub(r"(理论部分)(\d{1,3})(?=\s*\n|[一二三四五六七八九十百千])", r"\1\n\2", t)
    return t


def _classify_section_header_line(ln: str) -> Optional[str]:
    """识别章节行：一、单选题 / 二、多选题 等，返回题型。"""
    s = ln.strip()
    if not re.match(r"^[一二三四五六七八九十百千]+[、．.]", s):
        return None
    if "多选" in s or "多项选择" in s:
        return "multiple"
    if "单选" in s or "单项" in s:
        return "single"
    if "判断" in s:
        return "judge"
    if "填空" in s:
        return "fill"
    return None


def _extract_sections_from_block(block: str, section: Optional[str]) -> Tuple[str, Optional[str]]:
    """移除块内章节标题行，并更新当前章节题型（封面与「一、单选题」可同块）。"""
    lines = block.splitlines()
    sec = section
    kept: List[str] = []
    for ln in lines:
        hit = _classify_section_header_line(ln)
        if hit:
            sec = hit
            continue
        kept.append(ln)
    return "\n".join(kept).strip(), sec


def _strip_stem_cover_lines(stem: str) -> str:
    """去掉题干前封面、独立页码行、章节标题行等噪声。"""
    lines = stem.splitlines()
    out: List[str] = []
    started = False
    for ln in lines:
        st = ln.strip()
        if not started:
            if not st:
                continue
            if re.match(r"^生成式\s*AI\s*工程师", st):
                continue
            if re.match(r"^理论部分", st):
                continue
            if re.match(r"^\d{1,3}$", st):
                continue
            if _classify_section_header_line(st):
                continue
            started = True
        out.append(ln)
    return "\n".join(out).strip()


def _strip_leading_qnums(stem: str) -> str:
    """去掉题干行首题号（可多层，如 1. 或 1、）。"""
    s = stem.strip()
    while True:
        ns = re.sub(r"^\s*\d{1,3}[\.\)、]\s*", "", s, count=1)
        if ns == s:
            break
        s = ns
    return s.strip()


def _clean_answer_raw(s: str) -> str:
    """去掉答案文本尾部粘连的页码等，如「正确118」「B12」误判场景尽量少误伤。"""
    s = (s or "").strip()
    if not s:
        return s
    # 判断题常见：正确/错误 后粘连 1～3 位数字
    s = re.sub(r"(正确|错误|对|错)([。．]?\s*)(\d{1,3})\s*$", r"\1\2", s)
    # 选择题单个字母后粘连页码
    if re.match(r"^[A-Ea-e](\d{1,3})\s*$", s):
        s = re.sub(r"^([A-Ea-e])\d{1,3}\s*$", r"\1", s)
    return s.strip()


def _crop_analysis_body(raw: str) -> str:
    """截取「解析」正文，遇下一题题号或下一道答案行时截断，避免误吞整卷。"""
    if not raw:
        return ""
    lines_out: List[str] = []
    for line in raw.splitlines():
        st = line.strip()
        if lines_out:
            # 下一道题常见行首：数字+点/顿号+正文
            if re.match(r"^\d{1,3}[\.\)、]\s*\S", st):
                break
            # 下一题的「答案：」误入解析区
            if re.match(r"^(?:答案|标准答案)[：:]", st):
                break
        elif re.match(r"^\d{1,3}[\.\)、]\s*\S", st):
            # 解析区首行即像下一题题号，视为无有效解析
            return ""
        lines_out.append(line)
    return "\n".join(lines_out).strip()


def _tail_looks_like_promo_fragment(tail: str) -> bool:
    """判断从「更多」起的片段是否为广告/联系方式（用于截断，不误删题干）。"""
    u = tail.strip()
    if len(u) < 10:
        return False
    if re.search(r"1[3-9]\d{9}", u) and re.search(r"(?:微信|老师|联系|扫码)", u):
        return True
    if u.startswith("更多") and "联系" in u:
        return True
    if "更多" in u and re.search(r"(?:请联系|联系[：:]|微信)", u):
        return True
    return False


def _is_promotional_only_line(line: str) -> bool:
    """整行可视为纯广告/联系方式（整行删除，不与题干混在同一行末段时）。"""
    t = line.strip()
    if len(t) < 8:
        return False
    if t.startswith("更多"):
        return True
    if re.search(r"1[3-9]\d{9}", t) and re.search(r"(?:微信|老师|联系|扫码|公众号)", t):
        return True
    if re.search(r"请联系[：:｜|]", t) and len(t) < 500:
        return True
    return False


def _strip_mixed_line_trailing_ad(line: str) -> str:
    """一行内题干后粘连「更多…联系/手机…」时，只去掉广告尾段。"""
    if "更多" not in line:
        return line
    m = re.search(r"(更多.{8,600}?(?:请联系|联系[：:]|微信|1[3-9]\d{9}).*)$", line)
    if not m:
        return line
    if _tail_looks_like_promo_fragment(m.group(1)):
        return line[: m.start()].rstrip()
    return line


def _strip_inline_promo_after_sentence(s: str) -> str:
    """句读后同一字符串内粘连的广告尾（换行前）。"""
    if not s or "更多" not in s:
        return s
    m = re.search(
        r"([。！？])\s*(更多.{8,600}?(?:请联系|联系[：:]|微信|1[3-9]\d{9}).*)$",
        s,
        re.DOTALL,
    )
    if m and _tail_looks_like_promo_fragment(m.group(2)):
        return s[: m.start(2)].rstrip()
    return s


def _strip_irrelevant_import_tail(s: str) -> str:
    """去掉题干/选项/解析尾部与题库无关的广告、联系方式（先处理行内粘连，再删纯广告行）。"""
    if not s:
        return ""
    s = _strip_inline_promo_after_sentence(s)
    lines = s.splitlines()
    if lines:
        lines[-1] = _strip_mixed_line_trailing_ad(lines[-1])
    while lines:
        last = lines[-1].strip()
        if not last:
            lines.pop()
            continue
        if _is_promotional_only_line(lines[-1]):
            lines.pop()
            continue
        break
    out = "\n".join(lines).strip()
    return _strip_inline_promo_after_sentence(out)


def _strip_field_page_noise(s: str) -> str:
    """题干/选项/解析共用：去掉 PDF 粘连页码（整行数字、标点与问号后数字、括号后数字等）。"""
    if not s:
        return ""
    lines_out: List[str] = []
    for ln in s.splitlines():
        t = ln.rstrip()
        if not t:
            lines_out.append("")
            continue
        if re.match(r"^\d{1,3}\s*$", t):
            continue
        t = re.sub(r"([。．！!？?，,；;])\s*(\d{1,3})\s*$", r"\1", t)
        t = re.sub(r"^(.+[\u4e00-\u9fff。．！!？?，,；;])\s*(\d{1,3})\s*$", r"\1", t)
        t = re.sub(r"([？?！!])\s*\*{0,3}\s*(\d{1,3})\s*\*{0,3}\s*$", r"\1", t)
        t = re.sub(r"([？?！!])\s+(\d{1,3})\s*$", r"\1", t)
        t = re.sub(r"(\))\s*(\d{1,3})\s*$", r"\1", t)
        lines_out.append(t)
    return "\n".join(lines_out).rstrip()


def _strip_analysis_page_noise(s: str) -> str:
    """解析字段页码清理（与题干/选项规则一致）。"""
    return _strip_field_page_noise(s)


def _truncate_analysis(s: str) -> Optional[str]:
    """解析入库前去掉页码噪声、广告尾句并截断长度。"""
    s = _strip_analysis_page_noise(s)
    s = _strip_irrelevant_import_tail(s)
    s = s.strip()
    if not s:
        return None
    if len(s) > ANALYSIS_MAX_LEN:
        return s[:ANALYSIS_MAX_LEN]
    return s


def normalize_analysis(value: Optional[str]) -> Optional[str]:
    """供 API 入库前统一限制解析长度（与导入解析逻辑一致）。"""
    if value is None:
        return None
    return _truncate_analysis(value)


def _split_answer_analysis(text: str) -> Tuple[str, Optional[str], Optional[str]]:
    """从整段文本去掉「答案」「解析」行块，仅保留题目与选项部分。"""
    t = text.strip()
    aw: Optional[str] = None
    an: Optional[str] = None
    # 先去掉「解析：…」：不可贪婪到全文末尾，需按下一题边界裁剪后再截断长度
    m_p = re.search(r"(?:^|\n)\s*解析[：:]\s*", t)
    if m_p:
        rest = t[m_p.end() :]
        an = _truncate_analysis(_crop_analysis_body(rest))
        t = t[: m_p.start()].strip()
    # 再去掉「答案：」单行
    m_a = re.search(r"(?:^|\n)\s*(?:答案|标准答案)[：:]\s*([^\n]+)", t)
    if m_a:
        aw = _clean_answer_raw(m_a.group(1))
        t = t[: m_a.start()].strip()
    return t, aw, an


def _parse_answer_to_json(ans: Optional[str], q_type: str) -> Any:
    """根据文本答案生成 answer_json。"""
    if not ans or not str(ans).strip():
        if q_type == "multiple":
            return {"choices": ["A"]}
        if q_type == "judge":
            return {"choice": "T"}
        if q_type == "fill":
            return {"text": ""}
        return {"choice": "A"}
    s = _clean_answer_raw(str(ans))
    if q_type == "multiple":
        keys = re.findall(r"[A-E]", s.upper())
        return {"choices": keys if keys else ["A"]}
    if q_type == "judge":
        if any(x in s for x in ("错", "误", "否", "×")) or s in ("错误", "不正确", "F"):
            return {"choice": "F"}
        if any(x in s for x in ("对", "是", "正确", "√", "T", "真")):
            return {"choice": "T"}
        return {"choice": "T"}
    if q_type == "fill":
        return {"text": s}
    m = re.match(r"^([A-Ea-e])", s)
    if m:
        return {"choice": m.group(1).upper()}
    return {"choice": "A"}


def _parse_option_lines(lines: List[str]) -> Tuple[List[Dict[str, str]], List[str]]:
    """解析 A. B. C. 选项行。"""
    opts: List[Dict[str, str]] = []
    stem_lines: List[str] = []
    opt_pat = re.compile(r"^\s*([A-Ea-e])[\.\．、:：]\s*(.+)$")
    for ln in lines:
        m = opt_pat.match(ln.strip())
        if m:
            k = m.group(1).upper()
            opts.append({"key": k, "text": m.group(2).strip()})
        else:
            stem_lines.append(ln)
    return opts, stem_lines


def _infer_qtype(stem_core: str, opts: List[Dict[str, str]], hint: str) -> str:
    full = hint
    if re.search(r"判断题|对的打|错的打", full) and len(opts) <= 2:
        return "judge"
    if "多选" in full or "多项选择" in full:
        return "multiple"
    if "____" in full or "填空" in full or re.search(r"_{3,}|（\s{2,}", full):
        return "fill"
    if len(opts) >= 2:
        return "single"
    if re.search(r"正确|错误|对错", stem_core) and len(opts) < 2:
        return "judge"
    return "single"


def _is_noise_block(block: str) -> bool:
    """非题目说明、章节标题等不入库。"""
    t = block.strip()
    if len(t) < 8:
        return True
    # 含答案行的一般为有效题目（含判断题无 A/B 选项）
    if re.search(r"(?:答案|标准答案)[：:]", t):
        return False
    # 仅「一、单选题」类标题
    if re.match(r"^[一二三四五六七八九十百千]+[、．.].{0,40}$", t) and "？" not in t and "?" not in t:
        return True
    if re.match(r"^[（(]\s*[一二三四五六七八九十]+\s*[）)]", t) and len(t) < 60:
        return True
    # 无 A. B. 选项、无问号、无填空特征：整段视为封面/前言/目录，不保存
    if not re.search(r"[A-E][\.．、]\s*\S", t):
        if "？" not in t and "?" not in t and "____" not in t and "填空" not in t:
            return True
    if re.search(r"模拟题|理论部分|考试说明", t) and not re.search(r"[A-E][\.．、]", t):
        if "？" not in t and "?" not in t:
            return True
    return False


def _split_blocks(text: str) -> List[str]:
    t = text.strip()
    if not t:
        return []
    # 按空行或「行首题号」切题；题号后可为空格或直接跟中文（PDF 常见 2.以下）
    q_anchor = r"(?=\n\s*\d{1,3}[\.\)、])"
    parts = re.split(r"(?:\n\s*){2,}|" + q_anchor, t)
    out = [p.strip() for p in parts if p.strip()]
    if len(out) <= 1 and "\n" in t:
        parts2 = re.split(q_anchor, t)
        out = [p.strip() for p in parts2 if p.strip()]
    return out if out else [t]


def _apply_section_hint(
    stem_core: str,
    opts: List[Dict[str, str]],
    hint: str,
    section_hint: Optional[str],
) -> str:
    """结合卷面章节（一、单选题 / 二、多选题 等）确定题型。"""
    base = _infer_qtype(stem_core, opts, hint)
    has_ab_options = len(opts) >= 2 and bool(re.search(r"[A-E][\.．、]\s*\S", hint))
    if section_hint == "judge":
        # 判断题通常无 A/B 选项行；若仍有选项则按题干推断，避免章节状态错位把单选/多选判成判断
        if not has_ab_options:
            return "judge"
        return base
    if section_hint == "fill":
        return "fill"
    if section_hint == "multiple" and len(opts) >= 2:
        return "multiple"
    if section_hint == "single" and len(opts) >= 2:
        return "single"
    return base


def parse_question_block(block: str, section_hint: Optional[str] = None) -> Optional[Dict[str, Any]]:
    """解析单题：题干不含答案/解析；答案与解析分字段。"""
    if _is_noise_block(block):
        return None
    body, aw_raw, an_raw = _split_answer_analysis(block)
    lines = [ln for ln in body.splitlines() if ln.strip()]
    if not lines:
        return None
    opts, stem_lines = _parse_option_lines(lines)
    stem_core = "\n".join(stem_lines).strip()
    stem_core = _strip_stem_cover_lines(stem_core)
    stem_core = _strip_leading_qnums(stem_core)
    hint = body
    q_type = _apply_section_hint(stem_core, opts, hint, section_hint)

    if q_type == "fill":
        stem = _truncate_stem(stem_core if stem_core else body[:STEM_MAX_LEN])
        return {
            "q_type": "fill",
            "stem": stem,
            "options_json": None,
            "answer_json": _parse_answer_to_json(aw_raw, "fill"),
            "analysis": an_raw,
        }

    if q_type == "judge":
        stem = _truncate_stem(stem_core if stem_core else "\n".join(lines[:3]))
        opts_use = (
            opts
            if len(opts) >= 2
            else [{"key": "T", "text": "正确"}, {"key": "F", "text": "错误"}]
        )
        return {
            "q_type": "judge",
            "stem": stem,
            "options_json": opts_use,
            "answer_json": _parse_answer_to_json(aw_raw, "judge"),
            "analysis": an_raw,
        }

    stem = _truncate_stem(stem_core if stem_core else (lines[0] if lines else "（题干）"))
    if len(opts) < 2:
        # 无选项时退化为单选占位
        return {
            "q_type": "single",
            "stem": stem,
            "options_json": [
                {"key": "A", "text": "选项A"},
                {"key": "B", "text": "选项B"},
                {"key": "C", "text": "选项C"},
                {"key": "D", "text": "选项D"},
            ],
            "answer_json": _parse_answer_to_json(aw_raw, "single"),
            "analysis": an_raw,
        }

    ans_j = _parse_answer_to_json(aw_raw, q_type)
    return {
        "q_type": q_type,
        "stem": stem,
        "options_json": opts,
        "answer_json": ans_j,
        "analysis": an_raw,
    }


def finalize_import_item(item: Dict[str, Any]) -> Dict[str, Any]:
    """导入入库前再清一次页码噪声、广告尾句（题干、选项），解析走 normalize_analysis。"""
    stem = _strip_irrelevant_import_tail(_strip_field_page_noise((item.get("stem") or "").strip()))
    item["stem"] = _truncate_stem(stem)
    opts = item.get("options_json")
    if isinstance(opts, list):
        for o in opts:
            if isinstance(o, dict) and "text" in o:
                o["text"] = _strip_irrelevant_import_tail(
                    _strip_field_page_noise(str(o.get("text") or ""))
                )
    an = item.get("analysis")
    if an:
        item["analysis"] = normalize_analysis(str(an))
    else:
        item["analysis"] = None
    return item


def _parse_leading_question_number(block: str) -> Tuple[Optional[int], str]:
    """从块首提取题号（如 169. / 169、），便于多文件与答案册按编号对齐。"""
    b = block.lstrip()
    m = re.match(r"^(\d{1,4})[\.\．、）)]\s*", b)
    if not m:
        return None, block
    return int(m.group(1)), block


def _normalize_key_answer_line(s: str) -> str:
    s = (s or "").strip()
    if not s:
        return s
    trans = str.maketrans(
        "\uff21\uff22\uff23\uff24\uff25\uff41\uff42\uff43\uff44\uff45",
        "ABCDEabcde",
    )
    return s.translate(trans)


def _qtype_hint_from_answer_key(raw: str) -> str:
    """根据参考答案字母个数推断单选/多选；无字母时退回单选。"""
    u = re.sub(r"[\s,，;；]+", "", (raw or "").upper())
    keys = re.findall(r"[A-E]", u)
    if len(keys) > 1:
        return "multiple"
    return "single"


def _default_options_json_for_qtype(q_type: str) -> Any:
    if q_type == "judge":
        return [{"key": "T", "text": "正确"}, {"key": "F", "text": "错误"}]
    if q_type == "fill":
        return None
    return [
        {"key": "A", "text": "选项A"},
        {"key": "B", "text": "选项B"},
        {"key": "C", "text": "选项C"},
        {"key": "D", "text": "选项D"},
    ]


def _crop_analysis_after_key_tag(raw: str) -> str:
    """答案册中【试题解析】正文截断到下一题「数字.【参考答案】」之前。"""
    if not raw:
        return ""
    m = re.search(r"(?m)(?=^\d+\.\s*【参考答案】)", raw)
    if m:
        raw = raw[: m.start()]
    return raw.strip()


def parse_answer_key_file(text: str) -> Dict[int, Dict[str, Optional[str]]]:
    """解析单独答案册：「167. 【参考答案】 D」+「【试题解析】…」按题号索引。"""
    t = _normalize_import_text(text).strip()
    out: Dict[int, Dict[str, Optional[str]]] = {}
    if not t:
        return out
    for m in re.finditer(
        r"(?ms)^(\d+)\.\s*【参考答案】\s*([^\n\r]+?)(?:\s*\n\s*【试题解析】\s*([\s\S]*?))?(?=\n\s*\d+\.\s*【参考答案】|\Z)",
        t,
    ):
        num = int(m.group(1))
        ans_line = _normalize_key_answer_line(m.group(2))
        ans_body = re.sub(r"[\s　]+", "", ans_line).upper()
        an_raw = m.group(3)
        analysis: Optional[str] = None
        if an_raw is not None:
            body = _crop_analysis_after_key_tag(an_raw.strip())
            analysis = _truncate_analysis(body)
        out[num] = {"answer_raw": ans_body, "analysis": analysis}
    return out


def _looks_like_answer_key_file(text: str) -> bool:
    """是否为「仅参考答案+解析」类文件；与整卷含大量选项的混排区分。"""
    if not re.search(r"(?m)\d+\.\s*【参考答案】", text or ""):
        return False
    ref_cnt = len(re.findall(r"(?m)\d+\.\s*【参考答案】", text))
    opt_cnt = len(re.findall(r"(?m)^\s*[A-Ea-e][\.．、:：]\s*\S", text))
    if opt_cnt >= max(ref_cnt * 3, 15) and ref_cnt <= max(opt_cnt // 4, 8):
        return False
    return True


def _merge_parsed_items_for_num(store: Dict[int, Dict[str, Any]], num: int, item: Dict[str, Any]) -> None:
    """同一题号多份题干文件合并：取长题干与选项；解析/答案以非空补缺。"""
    if num not in store:
        store[num] = {**item}
        return
    o = store[num]
    if len((item.get("stem") or "")) > len((o.get("stem") or "")):
        o["stem"] = item["stem"]
        o["options_json"] = item.get("options_json")
        o["q_type"] = item.get("q_type") or o.get("q_type")
    if not (o.get("analysis") or "").strip() and (item.get("analysis") or "").strip():
        o["analysis"] = item["analysis"]


def _apply_answer_key_to_item(item: Dict[str, Any], ad: Dict[str, Optional[str]]) -> Dict[str, Any]:
    """将答案册中的答案、解析写入已解析题目（可按字母数修正多选）。"""
    raw = (ad.get("answer_raw") or "").strip()
    if not raw:
        if ad.get("analysis"):
            item["analysis"] = normalize_analysis(str(ad["analysis"]))
        return item
    qt = str(item.get("q_type") or "single")
    inferred = _qtype_hint_from_answer_key(raw)
    if inferred == "multiple" and qt == "single":
        item["q_type"] = "multiple"
        qt = "multiple"
    item["answer_json"] = _parse_answer_to_json(raw, qt)
    if ad.get("analysis"):
        item["analysis"] = normalize_analysis(str(ad["analysis"]))
    return item


def build_question_items_from_text(text: str) -> List[Tuple[Optional[int], Dict[str, Any]]]:
    """解析题目文件，返回 (题号或 None, 题目字典) 列表。"""
    t = _normalize_import_text(text)
    blocks = _split_blocks(t)
    out: List[Tuple[Optional[int], Dict[str, Any]]] = []
    section: Optional[str] = None
    for b in blocks:
        num, b0 = _parse_leading_question_number(b)
        b2, section = _extract_sections_from_block(b0, section)
        if not b2.strip():
            continue
        try:
            item = parse_question_block(b2, section_hint=section)
            if item:
                out.append((num, finalize_import_item(item)))
        except Exception:
            continue
    return out


def build_questions_from_uploaded_texts(sources: List[Tuple[str, str]]) -> Tuple[List[Dict[str, Any]], List[str]]:
    """多文件合并：题干类文件按题号合并；含「【参考答案】」的册子按题号对齐答案与解析。"""
    logs: List[str] = []
    by_num: Dict[int, Dict[str, Any]] = {}
    answers_by_num: Dict[int, Dict[str, Optional[str]]] = {}
    unnumbered: List[Dict[str, Any]] = []

    for fn, text in sources:
        label = fn or "upload"
        if not (text or "").strip():
            logs.append(f"[跳过] {label}：无文本内容")
            continue
        if _looks_like_answer_key_file(text):
            amap = parse_answer_key_file(text)
            logs.append(f"[解析] {label}：答案/解析册，题号条目 {len(amap)}")
            for k, v in amap.items():
                answers_by_num[k] = v
            continue
        pairs = build_question_items_from_text(text)
        logs.append(f"[解析] {label}：题目块 {len(pairs)}")
        for num, item in pairs:
            if num is None:
                unnumbered.append(item)
            else:
                _merge_parsed_items_for_num(by_num, num, item)

    for num in sorted(by_num.keys()):
        if num in answers_by_num:
            by_num[num] = finalize_import_item(_apply_answer_key_to_item(by_num[num], answers_by_num[num]))

    only_ans = set(answers_by_num.keys()) - set(by_num.keys())
    for num in sorted(only_ans):
        ad = answers_by_num[num]
        qt = _qtype_hint_from_answer_key(ad.get("answer_raw") or "")
        stub = {
            "q_type": qt,
            "stem": _truncate_stem(f"【题号{num}】题干未在本次上传文件中解析到，请手动编辑"),
            "options_json": _default_options_json_for_qtype(qt),
            "answer_json": _parse_answer_to_json(ad.get("answer_raw") or "", qt),
            "analysis": ad.get("analysis"),
        }
        by_num[num] = finalize_import_item(stub)
        logs.append(f"[补缺] 题号 {num}：仅有答案/解析，已生成题干占位")

    ordered_nums = sorted(by_num.keys())
    out: List[Dict[str, Any]] = [by_num[n] for n in ordered_nums]
    out.extend(unnumbered)
    return out, logs


def build_questions_from_text(text: str) -> List[Dict[str, Any]]:
    pairs = build_question_items_from_text(text)
    return [p[1] for p in pairs]


def build_image_placeholder(filename: str) -> Dict[str, Any]:
    """图片文件无文本时生成占位草稿，便于后续人工维护。"""
    return finalize_import_item(
        {
            "q_type": "single",
            "stem": f"【图片导入】{filename}：请编辑题干与选项",
            "options_json": [{"key": "A", "text": "选项A"}, {"key": "B", "text": "选项B"}],
            "answer_json": {"choice": "A"},
            "analysis": None,
        }
    )
